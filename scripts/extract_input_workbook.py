"""One-off: extract the input-only data of the OBE attainment software from the
course workbook (inputs and calculations mixed) into a clean upload workbook
(docs/PROJECT.md section 7) plus an extraction-notes file.

Usage: py scripts/extract_input_workbook.py [source.xlsx] [output.xlsx]

Reads values only (data_only). Never prints student names or identifiers.
Cell addresses are fixed for the CET3011B workbook; every block is guarded by an
anchor-label assert so a changed layout fails loudly instead of extracting junk.
"""
import glob
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
SRC = sys.argv[1] if len(sys.argv) > 1 else glob.glob(str(ROOT / "data/main file/*.xlsx"))[0]
OUT = Path(sys.argv[2] if len(sys.argv) > 2 else ROOT / "data/sample/CET3011B_input.xlsx")
NOTES = OUT.with_name(OUT.stem + "_notes.md")

PO_COLS = [f"PO{i}" for i in range(1, 13)] + [f"PSO{i}" for i in range(1, 4)]
COS = [f"CO{i}" for i in range(1, 6)]
notes, checks = [], []  # notes: choices/gaps; checks: expected-vs-found


def num(v):
    return int(v) if isinstance(v, (int, float)) and float(v).is_integer() else v


wb = openpyxl.load_workbook(SRC, data_only=True)
t1, t2, t3 = wb.worksheets[0], wb.worksheets[1], wb.worksheets[2]


def anchor(ws, addr, text):
    got = str(ws[addr].value).strip()
    assert got.startswith(text), f"{ws.title}!{addr}: expected '{text}', found '{got}'"


# 1 + 2. Course header and COs (Tab 1) -------------------------------------
anchor(t1, "B7", "Course:")
m = re.match(r"Course:\s*(.+?)\s*\((\w+)\)\s*$", t1["B7"].value)
course_name, course_code = m.group(1), m.group(2)
m = re.match(r"Class:\s*(.+?)\s*\(AY\s*([\d-]+)\)", t1["H7"].value)
klass, acad_year = m.group(1), m.group(2)
semester = t1["M7"].value.split(":")[1].strip()
program = "B.Tech CSE"
batch = None
notes.append("Program: the workbook only says 'Final Year Btech, CSE'; written as 'B.Tech CSE' as in docs/PROJECT.md section 7. 'class' keeps the workbook text.")
notes.append("Batch: not stated anywhere in the workbook; left blank (AY 2025-26 final year suggests 2022-26, but that is an inference, not extracted).")

cos = []
for r in range(10, 15):
    assert t1.cell(r, 2).value == COS[r - 10]
    cos.append((t1.cell(r, 2).value, str(t1.cell(r, 3).value).strip()))
notes.append("CO descriptions are copied exactly as typed in the workbook, including its spelling errors (e.g. 'sanner', 'internediate').")

# 3. Table 1 ticks, 7. Table 7 final mapping -------------------------------
def po_table(ws, head_row, first_row, kind):
    anchor(ws, f"B{head_row}", "CO")
    assert [ws.cell(head_row, 3 + i).value for i in range(15)] == PO_COLS
    out = []
    for i, co in enumerate(COS):
        r = first_row + i
        assert ws.cell(r, 2).value == co
        vals = []
        for c in range(15):
            v = ws.cell(r, 3 + c).value
            if kind == "tick":
                v = "√" if v == "√" else (0 if v == 0 else None)
            else:
                v = num(v)
                assert v in (0, 1, 2, 3), (co, PO_COLS[c], v)
            vals.append(v)
        out.append([co] + vals)
    return out


ticks = po_table(t1, 18, 19, "tick")
final_map = po_table(t1, 80, 81, "num")
blank_ticks = [f"{row[0]}/{PO_COLS[i]}" for row in ticks for i, v in enumerate(row[1:]) if v is None]
if blank_ticks:
    notes.append(
        f"Table 1 has {len(blank_ticks)} empty cells (no tick, no 0) in columns PO9, PSO1, PSO2 for CO1 to CO5, while Table 2 lists PO9 and PSO1 to PSO3 for every unit "
        "and Table 7 maps them 2 or 3 (open item O4). Left blank, not guessed; the upload rule reads blank as 0.")

# 4. Table 2 units ----------------------------------------------------------
anchor(t1, "B26", "Table 2")
units = []
for r in range(28, 33):
    unit_no = int(re.search(r"\d+", t1.cell(r, 2).value).group())
    units.append([unit_no, num(t1.cell(r, 3).value), str(t1.cell(r, 4).value).replace(" ", "")])
anchor(t1, "B33", "Total")

# 5. Table 3 labs -----------------------------------------------------------
anchor(t1, "B35", "Table 3")
labs = []
for r in range(37, 45):
    labs.append([num(t1.cell(r, 2).value), str(t1.cell(r, 3).value).strip(), num(t1.cell(r, 4).value),
                 str(t1.cell(r, 5).value).replace(" ", ""), None])
notes.append("Lab hours: Table 3 has no hours column, so lab_hours is blank for all 8 assignments (open item O5). "
             "Table 4 (calculated, not extracted) shows practical hours 6, 8, 16 for CO1 to CO3 and 30 in total; docs section 7 gives 3, 3, 4, 4, 4, 4, 4, 4. "
             "Neither is written here.")
notes.append("Assignment names are copied as typed (e.g. 'Compund statement').")

# 6. Tools and syllabus mark allocation (Tab 2) ------------------------------
anchor(t2, "A45", "Actual value wise")
assert [t2.cell(46, c).value for c in (2, 3, 4, 7)] == ["Assignment 1 (/10)", "Active Learning", "MID TERM (/15)", "LCA"]
alloc = {}  # tool -> [marks for CO1..CO5]
for tool, col in (("Assignment", 2), ("PBL", 3), ("Mid-term", 4), ("LCA", 7)):
    alloc[tool] = [num(t2.cell(47 + i, col).value) or 0 for i in range(5)]
notes.append("Tool names and maximums: Tab 2 caption 'Assignment 1 (/10)' and the CCA distribution (10, 5, 15) conflict with Tab 3 ('Assignment 1 (/5)', 'PBL(/10)'). "
             "Used assignment 5 and PBL (Active Learning in Tab 2) 10, as settled in faculty meeting 2 (docs section 13).")
notes.append("Marks per CO: empty cells of the Tab 2 allocation table are written as 0 (a tool with no marks for a CO, docs section 8 step 2; docs worked example shows the same 0s). "
             "The LCA values (7.5, 7.5, 15) are the workbook's own split of Table 3; they equal 30 x assignments per CO / 8.")
notes.append("MCQ (/10) and Attendance (/5) appear in the workbook headers but every student has 0 and Tab 2 allocates them 0 marks; they are not tools of this course "
             "(they would break the internal total of 60), so they are left out of Tools and Marks (open item O8).")
notes.append("End-term: the workbook gives no per-CO marks (it counts equally for every CO), so marks_for_CO1 to CO5 are blank for the End-term row.")
tools = [["Assignment", "CCA", 5], ["PBL", "CCA", 10], ["Mid-term", "CCA", 15], ["LCA", "LCA", 30], ["End-term", "END_TERM", 40]]
tools = [t + (alloc[t[0]] if t[0] in alloc else [None] * 5) for t in tools]

# 8. CO targets (Tab 3, Table A / Form 8) -----------------------------------
anchor(t3, "C994", "Target attainment")
targets = []
for i, co in enumerate(COS):
    assert t3.cell(995 + i, 10).value == co
    targets.append([co, num(t3.cell(995 + i, 3).value)])

# 9. Student marks (Tab 3 Step-1 + end-term of Step-3) -----------------------
# Step-1 repeats the raw tool marks inside each CO block; the CO1 block is used and
# the other four are checked to be identical. End-term (typed) is in the Step-3 table.
anchor(t3, "B11", "Step-1")
anchor(t3, "A656", "Step-3")
assert t3["H660"].value.startswith("End Term (40)")
marks, absent_blank, stray = [], [], 0
for r in range(14, 305):
    serial, name = num(t3.cell(r, 1).value), str(t3.cell(r, 2).value).strip()
    stray += name != str(t3.cell(r, 2).value)
    r3 = r + 647
    assert t3.cell(r3, 1).value == t3.cell(r, 1).value and str(t3.cell(r3, 2).value).strip() == name, "Step-1/Step-3 rows misaligned"
    for t in (0, 1, 2, 5):  # Assignment, PBL, Mid term, LCA
        vals = {t3.cell(r, 3 + t + 6 * k).value for k in range(5)}
        assert len(vals) == 1, f"CO blocks disagree for serial {serial}"
    row = [serial, name] + [num(t3.cell(r, 3 + t).value) for t in (0, 1, 2, 5)] + [num(t3.cell(r3, 8).value)]
    if None in row[2:]:
        absent_blank.append((serial, [h for h, v in zip(["Assignment", "PBL", "Mid-term", "LCA", "End-term"], row[2:]) if v is None]))
    marks.append(row)
if absent_blank:
    notes.append("Blank marks in the workbook, left blank: " + "; ".join(f"serial {s}: {', '.join(h)}" for s, h in absent_blank) + ".")
gaps = sorted(set(range(1, int(marks[-1][0]) + 1)) - {r[0] for r in marks})
notes.append(f"Student identifier: the workbook only has a serial number ({len(marks)} students, serials 1 to {marks[-1][0]} with gaps at {', '.join(map(str, gaps))}). "
             "Serial number is written as student_id; the real roll number / PRN is missing (open item O6) and must be added before loading.")
notes.append(f"Absent students: the workbook has no marker, so absent cannot be told from a real zero (open item O7). {stray} student names had stray leading or trailing spaces, which were trimmed.")
notes.append("Students sheet (template 1) is not produced: the student list comes from the Marks sheet (docs section 7, MVP).")

# write output ----------------------------------------------------------------
sheets = {
    "Course_COs": (["course_code", "course_name", "program", "class", "semester", "academic_year", "batch", "co_code", "co_description"],
                   [[course_code, course_name, program, klass, semester, acad_year, batch, c, d] for c, d in cos]),
    "Units": (["unit_no", "theory_hours", "co_codes"], units),
    "Labs": (["assignment_no", "name", "unit_no", "co_codes", "lab_hours"], labs),
    "CO_PO_Ticks": (["co_code"] + PO_COLS, ticks),
    "CO_PO_Final": (["co_code"] + PO_COLS, final_map),
    "Tools": (["tool_name", "category", "max_marks"] + [f"marks_for_{c}" for c in COS], tools),
    "CO_Targets": (["co_code", "target_pct"], targets),
    "Marks": (["student_id", "student_name", "Assignment", "PBL", "Mid-term", "LCA", "End-term"], marks),
}
out = openpyxl.Workbook()
out.remove(out.active)
counts = {}
for name, (head, rows) in sheets.items():
    ws = out.create_sheet(name)
    ws.append(head)
    for row in rows:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(len(str(c.value or "")) for c in col[:50]) + 2)
    counts[name] = len(rows)
OUT.parent.mkdir(parents=True, exist_ok=True)
out.save(OUT)

# checks against the documentation -----------------------------------------
def chk(label, expected, found):
    checks.append((label, expected, found, "OK" if expected == found else "MISMATCH"))


chk("Course code", "CET3011B", course_code)
chk("Number of COs", 5, len(cos))
chk("Lab assignments", 8, len(labs))
chk("Students", 291, len(marks))
chk("Theory hours total", 45, sum(u[1] for u in units))
checks.append(("Lab hours total", 30, "not in workbook (lab_hours blank)", "NOT VERIFIABLE"))
chk("Internal tool maximums", [5, 10, 15, 30], [t[2] for t in tools[:4]])
chk("Internal maximum total", 60, sum(t[2] for t in tools[:4]))
chk("End-term maximum", 40, tools[4][2])
chk("CO targets", [60, 50, 55, 50, 50], [t[1] for t in targets])
chk("Tool marks per CO add up to tool maximum", [5, 10, 15, 30], [sum(alloc[t]) for t in ("Assignment", "PBL", "Mid-term", "LCA")])
chk("Marks above tool maximum", 0, sum(1 for r in marks for v, mx in zip(r[2:], [5, 10, 15, 30, 40]) if v is not None and v > mx))

lines = ["# CET3011B input extraction notes", "",
         f"Source: `{Path(SRC).name}`  \nOutput: `{OUT.name}`", "",
         "## Sheets", "", "| Sheet | Rows (excluding header) |", "| --- | --- |"]
lines += [f"| {k} | {v} |" for k, v in counts.items()]
lines += ["", "## Checks against docs/PROJECT.md", "", "| Check | Expected | Found | Result |", "| --- | --- | --- | --- |"]
lines += [f"| {a} | {b} | {c} | {d} |" for a, b, c, d in checks]
lines += ["", "## Choices and gaps", ""] + [f"- {n}" for n in notes]
lines += ["", "## Not extracted (calculated or phase 2)", "",
          "Hours per CO and baseline strengths (Tables 4 to 6), tool weights in percent and the marks-per-CO totals, step 2 and step 3 student values, end-term converted to 100, "
          "class averages, levels, statuses, Tables B, E, F, target and actual PO rows (incl. #REF!), and the survey tab."]
NOTES.write_text("\n".join(lines) + "\n", encoding="utf-8")

print(f"Wrote {OUT} and {NOTES.name}")
for k, v in counts.items():
    print(f"  {k}: {v} rows")
bad = [c for c in checks if c[3] != "OK"]
print("Checks:", "all OK" if not bad else "; ".join(f"{a} -> {d}" for a, _, _, d in bad))
