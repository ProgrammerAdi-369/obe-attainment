"""One-off: write the marks upload file from the Marks sheet of the extracted sample workbook.

    py scripts/extract_marks_file.py [input.xlsx] [marks.xlsx]

The input is data/sample/CET3011B_input.xlsx, made from the original course workbook by
scripts/extract_input_workbook.py. Output: data/sample/CET3011B_marks.xlsx, one sheet, one row per
student (student_id, student_name, Assignment, PBL, Mid-term, LCA, End-term). It holds real names,
so it stays in the git-ignored data/sample/ folder. Prints counts only, never names.
"""

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/sample/CET3011B_input.xlsx"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "data/sample/CET3011B_marks.xlsx"

rows = list(load_workbook(SRC, data_only=True)["Marks"].iter_rows(values_only=True))
wb = Workbook()
ws = wb.active
ws.title = "Marks"
for r in rows:
    ws.append(list(r))
wb.save(OUT)
print(f"wrote {OUT.name}: {len(rows) - 1} students, columns {list(rows[0])}")
