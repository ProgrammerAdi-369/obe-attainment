"""Marks ingest (docs/PROJECT.md section 7, template 8): read and validate the marks file.

One row per student: student_id, student_name, then one column per tool of the course. A bad file
gives row-level errors and loads nothing. A blank mark is not an error: it stays None, and the core
(score_students) flags that student and leaves them out of the processing.

Nothing here logs or prints. Errors carry row and column only, never a student id or name.
"""

import csv
from pathlib import Path
from typing import NamedTuple

from openpyxl import load_workbook

from obe.config import CourseConfig

ID_COLUMN, NAME_COLUMN = "student_id", "student_name"


class RowError(NamedTuple):
    row: int  # spreadsheet row number, the header is row 1
    column: str | None
    message: str


class MarksFileError(Exception):
    def __init__(self, errors: list[RowError]):
        self.errors = errors
        lines = [f"row {e.row}" + (f", {e.column}" if e.column else "") + f": {e.message}" for e in errors]
        super().__init__(f"marks file rejected, nothing loaded ({len(errors)} errors):\n  " + "\n  ".join(lines))


class MarksData(NamedTuple):
    students: dict[str, tuple[dict[str, float | None], float | None]]  # id -> (marks, end-term)
    names: dict[str, str]  # id -> name; personal data, keep out of logs


def read_table(path: Path, sheet: str | None = None) -> list[list]:
    """Rows of the file as lists; .xlsx (the named sheet, default the first) or .csv."""
    if Path(path).suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [[c if c != "" else None for c in row] for row in csv.reader(f)]
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def parse_marks(table: list[list], course: CourseConfig) -> MarksData:
    """Validate the rows against the course tools; raises MarksFileError listing every problem."""
    tools = {t.name: t for t in course.tools}
    end_term = next((t.name for t in course.tools if t.category == "END_TERM"), None)
    header = [_text(h) for h in table[0]] if table else []
    errors: list[RowError] = []

    for col in sorted({h for h in header if h and header.count(h) > 1}):
        errors.append(RowError(1, col, "duplicate column"))
    for col in header:
        if col and col not in tools and col not in (ID_COLUMN, NAME_COLUMN):
            errors.append(RowError(1, col, "unknown column"))
    for col in [ID_COLUMN, *tools]:
        if col not in header:
            errors.append(RowError(1, col, "missing column"))
    if errors:
        raise MarksFileError(errors)

    students: dict[str, tuple[dict[str, float | None], float | None]] = {}
    names: dict[str, str] = {}
    first_row: dict[str, int] = {}
    for n, cells in enumerate(table[1:], start=2):
        cells = list(cells) + [None] * (len(header) - len(cells))
        if all(_text(c) is None for c in cells):
            continue  # empty row
        row = dict(zip(header, cells))
        sid = _text(row[ID_COLUMN])
        if sid is None:
            errors.append(RowError(n, ID_COLUMN, "student id is blank"))
        elif sid in first_row:
            errors.append(RowError(n, ID_COLUMN, f"duplicate student id (first seen in row {first_row[sid]})"))
        else:
            first_row[sid] = n
        marks: dict[str, float | None] = {}
        for name, tool in tools.items():
            v = row[name]
            if isinstance(v, str):  # csv cells arrive as text
                try:
                    v = float(v)
                except ValueError:
                    pass
            if _text(v) is None:
                marks[name] = None  # blank: flagged later, not an error
            elif isinstance(v, bool) or not isinstance(v, (int, float)):
                errors.append(RowError(n, name, "mark is not a number"))
            elif not 0 <= v <= tool.max_marks:
                errors.append(RowError(n, name, f"mark is outside 0 to {tool.max_marks:g}"))
            else:
                marks[name] = float(v)
        if sid is not None and sid in first_row and first_row[sid] == n:
            end = marks.pop(end_term, None) if end_term else None
            students[sid] = (marks, end)
            if _text(row.get(NAME_COLUMN)):
                names[sid] = _text(row[NAME_COLUMN])
    if errors:
        raise MarksFileError(errors)
    return MarksData(students, names)


def load_marks(path: Path, course: CourseConfig, sheet: str | None = None) -> MarksData:
    return parse_marks(read_table(path, sheet), course)
