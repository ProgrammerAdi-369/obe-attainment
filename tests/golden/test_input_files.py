"""The extracted input workbook and marks file against the course YAML and each other.

Both files hold real names and are git-ignored, so these tests skip when they are absent.
"""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from obe.config import CONFIG_DIR, load_course, load_program_outcomes, load_rules
from obe.ingest.marks import load_marks

SAMPLE = Path(__file__).resolve().parents[2] / "data" / "sample"
INPUT, MARKS = SAMPLE / "CET3011B_input.xlsx", SAMPLE / "CET3011B_marks.xlsx"

pytestmark = pytest.mark.skipif(
    not (INPUT.exists() and MARKS.exists()), reason="sample files are not in git"
)


@pytest.fixture(scope="module")
def course():
    return load_course(CONFIG_DIR / "courses" / "CET3011B.yaml", load_rules(), load_program_outcomes())


@pytest.fixture(scope="module")
def sheets():
    wb = load_workbook(INPUT, read_only=True, data_only=True)
    try:
        # the Marks sheet is skipped here: it has the names and is checked through load_marks
        return {
            ws.title: [list(r) for r in ws.iter_rows(values_only=True)]
            for ws in wb.worksheets
            if ws.title != "Marks"
        }
    finally:
        wb.close()


def records(sheets, name):
    header, *rows = sheets[name]
    return [dict(zip(header, r)) for r in rows]


def codes(text):
    return [c.strip() for c in str(text).split(",")]


def test_sheets_are_the_upload_templates(sheets):
    assert set(sheets) == {
        "Course_COs", "Units", "Labs", "CO_PO_Ticks", "CO_PO_Final", "Tools", "CO_Targets",
    }


def test_header_and_cos_equal_the_yaml(sheets, course):
    rows = records(sheets, "Course_COs")
    h = course.course
    for r in rows:
        assert (r["course_code"], r["course_name"], r["program"], r["class"], r["semester"],
                r["academic_year"]) == (h.code, h.name, h.program, h.class_, h.semester, h.academic_year)
    assert [(r["co_code"], r["co_description"]) for r in rows] == [(c.code, c.description) for c in course.cos]
    assert rows[0]["batch"] is None and h.batch == "2022-26"  # the workbook has no batch; docs section 7


def test_units_and_labs_equal_the_yaml(sheets, course):
    units = [(int(r["unit_no"]), r["theory_hours"], codes(r["co_codes"])) for r in records(sheets, "Units")]
    assert units == [(u.unit_no, u.theory_hours, u.co_codes) for u in course.units]
    labs = records(sheets, "Labs")
    assert [(int(r["assignment_no"]), r["name"], int(r["unit_no"]), codes(r["co_codes"])) for r in labs] == [
        (x.assignment_no, x.name, x.unit_no, x.co_codes) for x in course.labs
    ]
    # O5: the workbook has no lab hours; the YAML holds the working answer
    assert all(r["lab_hours"] is None for r in labs)
    assert [x.lab_hours for x in course.labs] == [3, 3, 4, 4, 4, 4, 4, 4]


def test_tools_equal_the_yaml_and_lca_equals_the_generated_split(sheets, course):
    alloc = course.allocation()
    cos = [c.code for c in course.cos]
    rows = records(sheets, "Tools")
    assert [(r["tool_name"], r["category"], r["max_marks"]) for r in rows] == [
        (t.name, t.category, t.max_marks) for t in course.tools
    ]
    for r in rows:
        per_co = [r[f"marks_for_{co}"] for co in cos]
        if r["category"] == "END_TERM":
            assert per_co == [None] * 5  # equal for all COs, no per-CO marks
        else:
            assert per_co == [alloc[r["tool_name"]][co] for co in cos]  # LCA: 7.5, 7.5, 15, 0, 0


def test_mappings_and_targets_equal_the_yaml(sheets, course):
    tick = lambda v: 1 if v == "√" else 0 if v in (None, 0) else v
    for name, table, conv in (
        ("CO_PO_Ticks", course.co_po_ticks, tick),
        ("CO_PO_Final", course.co_po_final, lambda v: v),
    ):
        for r in records(sheets, name):
            assert {p: conv(v) for p, v in r.items() if p != "co_code"} == table[r["co_code"]], name
    assert {r["co_code"]: r["target_pct"] for r in records(sheets, "CO_Targets")} == course.co_targets


def test_marks_sheet_of_the_input_file_equals_the_marks_file(course):
    from_input = load_marks(INPUT, course, sheet="Marks")
    from_marks = load_marks(MARKS, course)
    assert len(from_input.students) == len(from_marks.students) == 291
    assert from_input.students == from_marks.students
    assert from_input.names == from_marks.names
