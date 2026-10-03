"""The master Excel report for the sample course, from the real marks file (git-ignored)."""

import pytest
import xlsxwriter
from compare import MARKS, available
from openpyxl import load_workbook

from obe.config import CONFIG_DIR, load_course, load_program_outcomes, load_rules
from obe.ingest.marks import load_marks
from obe.pipeline import calculate
from obe.reports.excel import write_master_workbook


def test_the_formula_checker_catches_a_wrong_stored_value(tmp_path, check_formulas):
    wb = xlsxwriter.Workbook(str(tmp_path / "x.xlsx"))
    ws = wb.add_worksheet("S")
    ws.write_formula("A1", "=1+1", None, 3)  # stored 3, really 2
    ws.write_formula("A2", "=A1*2", None, 4)
    wb.close()
    n, bad = check_formulas(tmp_path / "x.xlsx")
    assert n == 2 and len(bad) == 1 and bad[0].startswith("S!A1")


@pytest.mark.skipif(not available(), reason="sample marks file is not in git")
def test_sample_report(tmp_path, check_formulas):
    rules, pos = load_rules(), load_program_outcomes()
    course = load_course(CONFIG_DIR / "courses" / "CET3011B.yaml", rules, pos)
    marks = load_marks(MARKS, course)
    result = calculate(course, rules, pos, marks)
    path = tmp_path / "report.xlsx"
    write_master_workbook(result, path)

    n, bad = check_formulas(path)
    assert bad == [] and n == 71  # 10 class averages, Table B (15), Table E (1), Table F (45)

    wb = load_workbook(path, data_only=True)
    steps = wb["Student steps"]
    assert sum(1 for r in steps.iter_rows(min_row=4, max_row=291) if r[0].value) == 288
    assert [wb["Table B"][f"C{r}"].value for r in range(4, 9)] == [2, 3, 3, 3, 3]
    assert wb["Table E"]["B3"].value == pytest.approx(2.8)
    assert wb["Table F"]["B9"].value == 3 and wb["Table F"]["I9"].value == pytest.approx(2.6)

    # names are personal data: none of them may be anywhere in the file
    texts = " ".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value)
    assert sum(1 for name in marks.names.values() if name in texts) == 0
