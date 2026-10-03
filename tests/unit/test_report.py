"""Master Excel report: sheet list, formulas, values, privacy. Synthetic students, real course YAML."""

import pytest
from openpyxl import load_workbook

from obe.cli import main
from obe.config import CONFIG_DIR, load_course, load_program_outcomes, load_rules
from obe.ingest.marks import MarksData
from obe.pipeline import calculate
from obe.reports.excel import write_master_workbook

S1 = {"Assignment": 3, "PBL": 3, "Mid-term": 6, "LCA": 12}  # the worked example, end-term 22
MARKS = MarksData(
    students={
        "1": (dict(S1), 22),
        "2": ({"Assignment": 5, "PBL": 9, "Mid-term": 14, "LCA": 28}, 36),
        "3": ({**S1, "Assignment": None}, 22),  # blank mark: flagged, left out
    },
    names={},
)


@pytest.fixture(scope="module")
def cfg():
    rules, pos = load_rules(), load_program_outcomes()
    return rules, pos, load_course(CONFIG_DIR / "courses" / "CET3011B.yaml", rules, pos)


@pytest.fixture(scope="module")
def run(cfg):
    rules, pos, course = cfg
    return calculate(course, rules, pos, MARKS)


@pytest.fixture(scope="module")
def book(run, tmp_path_factory):
    path = tmp_path_factory.mktemp("rep") / "r.xlsx"
    write_master_workbook(run, path)
    return path


def test_sheets_follow_section_11_with_c_d_and_student_po_left_out(book):
    assert load_workbook(book).sheetnames == [
        "Summary", "CO-PO mapping", "Tool weights", "Student steps", "Table A", "Table B",
        "Table E", "Table F", "Config",
    ]


def test_every_formula_recomputes_to_the_stored_value(book, check_formulas):
    n, bad = check_formulas(book)
    assert n > 40 and bad == []  # class averages, Table B, Table E, Table F


def test_formulas_are_where_the_spec_wants_them(book):
    wb = load_workbook(book)
    assert wb["Student steps"]["H7"].value == "=AVERAGE(H4:H5)"  # two scored students, rows 4 and 5
    assert wb["Table B"]["C4"].value.startswith("=IF(")
    assert wb["Table E"]["B3"].value == "=AVERAGE('Table B'!C4:C8)"
    assert wb["Table F"]["B9"].value.startswith("=IFERROR(AVERAGEIF(")


def test_tool_weights_sheet_shows_percentages_only(book):
    ws = load_workbook(book)["Tool weights"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5, max_row=9)]
    assert ws["A4"].value == "Tool"
    assert [r[0] for r in rows] == ["Assignment", "PBL", "Mid-term", "LCA", "Total"]
    assert [round(x, 2) for x in rows[0][1:]] == [17.14, 11.43, 0, 0, 0]  # not the marks 3 and 2
    assert all(r[j] <= 100 for r in rows for j in range(1, 6))
    assert [round(x) for x in rows[4][1:]] == [100] * 5


def test_stored_values_match_the_calculation(book, run):
    wb = load_workbook(book, data_only=True)
    b = wb["Table B"]
    assert [b[f"C{r}"].value for r in range(4, 9)] == [x.final_level for x in run.classes.values()]
    assert b["D4"].value == run.classes["CO1"].status
    assert wb["Table E"]["B3"].value == pytest.approx(run.overall)
    f = wb["Table F"]
    assert f["B9"].value == 3 and f["I9"].value == pytest.approx(2.6)  # PO1, PO8 targets
    assert f["E9"].value in (None, "")  # PO4: no CO maps to it
    assert f["B10"].value == pytest.approx(3 * run.overall / 3)


def test_student_rows_scored_and_flagged_use_ids_only(book):
    ws = load_workbook(book)["Student steps"]
    ids = [r[0].value for r in ws.iter_rows(min_row=4, max_row=5)]
    assert ids == ["1", "2"]
    text = [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)]
    assert any("Not scored: 1 students" in t for t in text) and "Assignment" in text


def test_table_b_has_empty_reason_and_action_plan_columns(book):
    ws = load_workbook(book)["Table B"]
    assert [ws["E3"].value, ws["F3"].value] == ["Reason for gap", "Action plan"]
    assert all(ws[f"{c}{r}"].value is None for c in "EF" for r in range(4, 9))


def test_student_po_sheet_appears_only_when_the_flag_is_on(cfg, tmp_path):
    rules, pos, course = cfg
    on = calculate(course, rules.model_copy(update={"student_po_enabled": True}), pos, MARKS)
    write_master_workbook(on, tmp_path / "on.xlsx")
    wb = load_workbook(tmp_path / "on.xlsx", data_only=True)
    assert "Student PO" in wb.sheetnames
    assert [wb["Student PO"][f"A{r}"].value for r in (4, 5)] == ["1", "2"]


def test_cli_writes_the_report_and_prints_counts_only(tmp_path, capsys):
    marks = tmp_path / "m.csv"
    marks.write_text(
        "student_id,student_name,Assignment,PBL,Mid-term,LCA,End-term\n"
        "SECRET-ID,Secret Name,3,3,6,12,22\nX2,Other Name,5,9,14,28,36\n", encoding="utf-8"
    )
    assert main(["report", "--course", "CET3011B", "--marks", str(marks), "--out", str(tmp_path / "o.xlsx")]) == 0
    out = capsys.readouterr().out
    assert "2 students scored, 0 flagged" in out and "SECRET" not in out and "Name" not in out
    cells = [c.value for ws in load_workbook(tmp_path / "o.xlsx").worksheets for r in ws.iter_rows() for c in r]
    assert "Secret Name" not in cells and "Other Name" not in cells


def test_cli_rejects_a_bad_file_and_writes_nothing(tmp_path, capsys):
    marks = tmp_path / "m.csv"
    marks.write_text("student_id,Assignment,PBL,Mid-term,LCA,End-term\n1,99,3,6,12,22\n", encoding="utf-8")
    out = tmp_path / "o.xlsx"
    assert main(["report", "--course", "CET3011B", "--marks", str(marks), "--out", str(out)]) == 1
    assert "row 2, Assignment" in capsys.readouterr().err and not out.exists()
